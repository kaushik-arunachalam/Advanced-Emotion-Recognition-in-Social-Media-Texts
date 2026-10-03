# Builds Capstone_Phase3_Slides.pptx in the exact visual style of Capstone_Phase2_Slides.pptx
# (16:9, white, Cambria titles, Calibri body, grey rules, F2F2F2 table headers, thin black process boxes).
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

BLACK, DARK, GREY, RULE, HDR, ARROW, WHITE = '000000', '404040', '6E6E6E', 'D9D9D9', 'F2F2F2', 'BFBFBF', 'FFFFFF'
X0, W = 0.70, 11.93
COL2 = 6.92; COLW = 5.71; DIV = 6.67

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
RGB = RGBColor.from_string
PAGE = [1]


# ---------------------------------------------------------------- primitives
def _box(slide, x, y, w, h, anchor='t'):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    return s, tf


def _run(p, text, font='Calibri', size=14, bold=False, italic=False, color=DARK, spc=None):
    r = p.add_run(); r.text = text
    f = r.font; f.name = font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = RGB(color)
    if spc: r._r.get_or_add_rPr().set('spc', str(spc))
    return r


def _para_fmt(p, line=None, after=None, bullet=False, align=None):
    pPr = p._p.get_or_add_pPr()
    if bullet:
        pPr.set('marL', '177800'); pPr.set('indent', '-177800')
    for tag in ('a:lnSpc', 'a:spcAft', 'a:buSzPct', 'a:buChar', 'a:buNone'):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    if line:
        e = etree.SubElement(pPr, qn('a:lnSpc')); etree.SubElement(e, qn('a:spcPct')).set('val', str(int(line * 1000)))
    if after is not None:
        e = etree.SubElement(pPr, qn('a:spcAft')); etree.SubElement(e, qn('a:spcPts')).set('val', str(int(after * 100)))
    if bullet:
        etree.SubElement(pPr, qn('a:buSzPct')).set('val', '100000'); etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
    if align: p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]


def text(slide, x, y, w, h, paras, font='Calibri', size=14, bold=False, italic=False, color=DARK, align='l', spc=None,
         anchor='t', line=None, after=None):
    """paras: str (\\n = new paragraph) or list of paragraphs; a paragraph is str or list of (text, overrides)."""
    s, tf = _box(slide, x, y, w, h, anchor)
    if isinstance(paras, str): paras = paras.split('\n')
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            _run(p, t, o.get('font', font), o.get('size', size), o.get('bold', bold), o.get('italic', italic), o.get('color', color), o.get('spc', spc))
        _para_fmt(p, line, after, align=align)
    return s


def bullets(slide, x, y, w, h, items, size=14, color=DARK, after=10, line=108):
    s, tf = _box(slide, x, y, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs = [(it, {})] if isinstance(it, str) else it
        for t, o in runs:
            _run(p, t, 'Calibri', o.get('size', size), o.get('bold', False), o.get('italic', False), o.get('color', color))
        _para_fmt(p, line, after, bullet=True)
    return s


def line(slide, x1, y1, x2, y2, color=RULE, width=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGB(color); c.line.width = Pt(width)
    st = c._element.find(qn('p:style'))
    if st is not None: c._element.remove(st)
    return c


def rect(slide, x, y, w, h, line_color=BLACK, fill=WHITE, width=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGB(fill)
    if line_color: s.line.color.rgb = RGB(line_color); s.line.width = Pt(width)
    else: s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ''
    return s


def arrow(slide, x, y, w=0.66, h=0.28):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGB(ARROW); s.line.fill.background(); s.shadow.inherit = False
    return s


def _cell_border(cell, color=RULE, w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for e in tcPr.findall(qn(tag)): tcPr.remove(e)
    for i, tag in enumerate(('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')):
        ln = etree.Element(qn(tag), w=str(w), cap='flat', cmpd='sng', algn='ctr')
        sf = etree.SubElement(ln, qn('a:solidFill')); etree.SubElement(sf, qn('a:srgbClr')).set('val', color)
        etree.SubElement(ln, qn('a:prstDash')).set('val', 'solid')
        tcPr.insert(i, ln)


def table(slide, x, y, col_w, rows, row_h, hdr_size=12.5, body_size=11.5, bold_first_col=False, col_align=None):
    """rows[0] = header. A cell is str ('\\n' = new paragraph) or list of paragraphs of (text, overrides) runs."""
    shp = slide.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(sum(col_w)), Inches(sum(row_h)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for a in list(tblPr.attrib): del tblPr.attrib[a]
    for e in list(tblPr): tblPr.remove(e)
    for j, cw in enumerate(col_w): tbl.columns[j].width = Inches(cw)
    for i, rh in enumerate(row_h): tbl.rows[i].height = Inches(rh)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = Emu(76200); c.margin_top = c.margin_bottom = Emu(50800)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid(); c.fill.fore_color.rgb = RGB(HDR if i == 0 else WHITE)
            tf = c.text_frame; tf.word_wrap = True
            paras = val.split('\n') if isinstance(val, str) else val
            for k, para in enumerate(paras):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                runs = [(para, {})] if isinstance(para, str) else para
                for t, o in runs:
                    hdr = i == 0
                    _run(p, t, 'Calibri', o.get('size', hdr_size if hdr else body_size),
                         o.get('bold', hdr or (bold_first_col and j == 0)), o.get('italic', False), o.get('color', BLACK if hdr else DARK))
                _para_fmt(p, align=(col_align[j] if col_align else 'l'))
            _cell_border(c)
    return shp


def page(slide):
    PAGE[0] += 1
    text(slide, 12.13, 7.02, 0.50, 0.25, str(PAGE[0]), size=10, color=GREY, align='r')


def title(slide, t, sub=None, sub_bold_prefix=None):
    text(slide, X0, 0.55, W, 0.70, t, font='Cambria', size=30, bold=True, color=BLACK)
    if sub:
        runs = ([(sub_bold_prefix, {'bold': True, 'color': BLACK})] if sub_bold_prefix else []) + [(sub, {})]
        text(slide, X0, 1.30, W, 0.40, [runs], size=13, italic=True, color=GREY if not sub_bold_prefix else DARK)


def heading(slide, x, y, w, t, size=16):
    text(slide, x, y, w, 0.40, t, font='Cambria', size=size, bold=True, color=BLACK)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def new_slide():
    s = prs.slides.add_slide(BLANK)
    return s


def chart_style(ch, legend=True, size=10):
    ch.font.name = 'Calibri'; ch.font.size = Pt(size); ch.font.color.rgb = RGB(DARK)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size); ch.legend.font.name = 'Calibri'
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = RGB('E6E6E6'); va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background(); va.tick_labels.font.size = Pt(size); va.tick_labels.font.color.rgb = RGB(GREY)
    ca.format.line.color.rgb = RGB(RULE); ca.tick_labels.font.size = Pt(size); ca.tick_labels.font.color.rgb = RGB(DARK)
    ca.has_major_gridlines = False


